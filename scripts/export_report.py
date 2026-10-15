"""Export BAO_CAO.md to Word/PDF and validate Vietnamese text and images."""
import json
import hashlib
import os
import textwrap
from html import escape
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from markdown_it import MarkdownIt
from PIL import Image as PILImage, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image, KeepTogether, PageBreak, Paragraph, Preformatted, SimpleDocTemplate,
    Spacer, Table, TableStyle,
)


OUT = Path("submission")
NAME = "Bao_cao_DAY27_MaiQuangDung_2A202602966"
FONTS = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"


def architecture():
    canvas = PILImage.new("RGB", (1600, 1000), "white")
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype(str(FONTS / "arial.ttf"), 25)
    bold = ImageFont.truetype(str(FONTS / "arialbd.ttf"), 30)

    def box(x, y, title, detail, fill="#eff4f8"):
        draw.rounded_rectangle((x-210, y-48, x+210, y+48), radius=12, fill=fill, outline="#8a9ba8", width=2)
        draw.text((x, y-18), title, fill="#17364a", font=bold, anchor="mm")
        draw.text((x, y+21), detail, fill="#334b5b", font=font, anchor="mm")

    def arrow(points):
        draw.line(points, fill="#526c7b", width=4)
        x, y = points[-1]
        px, py = points[-2]
        if y > py:
            tip = [(x, y), (x-8, y-14), (x+8, y-14)]
        elif x < px:
            tip = [(x, y), (x+14, y-8), (x+14, y+8)]
        else:
            tip = [(x, y), (x-14, y-8), (x-14, y+8)]
        draw.polygon(tip, fill="#526c7b")

    draw.text((800, 30), "Luồng xử lý HITL trên giao diện Streamlit", font=bold, fill="#17364a", anchor="mm")
    box(800, 120, "fetch_pr → analyze", "PR diff → PRAnalysis")
    box(800, 260, "route", "Đọc confidence")
    arrow([(800,168),(800,212)])
    for x in (270,800,1330):
        arrow([(800,308),(800,335),(x,335),(x,372)])
    box(270,420,"auto_approve","confidence ≥ 0.73", "#e4f3e8")
    box(800,420,"human_approval","0.58 ≤ confidence < 0.73", "#fff4d9")
    box(1330,420,"escalate","confidence < 0.58", "#fde8e4")
    box(270,600,"Đăng review","Không cần người duyệt", "#e4f3e8")
    box(800,600,"interrupt()","Approve / Reject / Edit", "#fff4d9")
    box(1330,600,"interrupt()","Câu hỏi → trả lời", "#fde8e4")
    for x in (270,800,1330):
        arrow([(x,468),(x,552)])
    box(1330,780,"synthesize","Tổng hợp review mới")
    arrow([(1330,648),(1330,732)])
    box(800,780,"commit / kết thúc","Đăng nếu được duyệt")
    arrow([(800,648),(800,732)])
    arrow([(1330,828),(1330,880),(1050,880),(1050,600),(1010,600)])
    draw.text((1250,925), "Web: duyệt cuối sau Q&A", font=font, fill="#334b5b", anchor="mm")
    draw.text((800,975), "Mọi bước có audit; checkpoint lưu trạng thái để resume cùng thread_id.", font=font, fill="#334b5b", anchor="mm")
    Path("artifacts").mkdir(exist_ok=True)
    canvas.save("artifacts/architecture.png")


def inline_pdf(token):
    parts = []
    link_tags = []
    for item in token.children or []:
        if item.type == "text": parts.append(escape(item.content))
        elif item.type == "code_inline": parts.append(f'<font face="Code">{escape(item.content)}</font>')
        elif item.type == "strong_open": parts.append("<b>")
        elif item.type == "strong_close": parts.append("</b>")
        elif item.type == "em_open": parts.append("<i>")
        elif item.type == "em_close": parts.append("</i>")
        elif item.type in {"softbreak", "hardbreak"}: parts.append("<br/>")
        elif item.type == "link_open":
            href = item.attrGet("href") or ""
            remote = href.startswith("http")
            parts.append(f'<a href="{escape(href, quote=True)}" color="#17577a">' if remote else "<font>")
            link_tags.append("</a>" if remote else "</font>")
        elif item.type == "link_close": parts.append(link_tags.pop())
    return "".join(parts)


def plain(token):
    return "".join(t.content if t.type in {"text", "code_inline"} else
                   " " if t.type in {"softbreak", "hardbreak"} else ""
                   for t in token.children or [])


def main():
    architecture()
    OUT.mkdir(exist_ok=True)
    for name, filename in [("Report", "times.ttf"), ("Report-Bold", "timesbd.ttf"),
                           ("Report-Italic", "timesi.ttf"), ("Report-BoldItalic", "timesbi.ttf"),
                           ("Code", "consola.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(FONTS / filename)))
    pdfmetrics.registerFontFamily("Report", normal="Report", bold="Report-Bold",
                                  italic="Report-Italic", boldItalic="Report-BoldItalic")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("BodyVI", fontName="Report", fontSize=11, leading=16, spaceAfter=7))
    styles.add(ParagraphStyle("TableVI", fontName="Report", fontSize=9, leading=12, spaceAfter=0))
    styles.add(ParagraphStyle("CaptionVI", fontName="Report-Italic", fontSize=9.5, leading=13, spaceAfter=8))
    styles.add(ParagraphStyle("TitleVI", fontName="Report-Bold", fontSize=24, leading=33, alignment=TA_CENTER, spaceAfter=22))
    for level in (1,2,3):
        styles.add(ParagraphStyle(f"HeadingVI{level}", fontName="Report-Bold", fontSize={1:18,2:15,3:12}[level],
                                  leading={1:23,2:20,3:17}[level], spaceBefore=13, spaceAfter=8, keepWithNext=True))
    code_style = ParagraphStyle("CodeVI", fontName="Code", fontSize=8, leading=11,
                                backColor=colors.HexColor("#f1f4f6"), borderPadding=7, spaceAfter=9)

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.text = "Mai Quang Dũng | Day27 – Track 3 | Trang "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.runs[0].font.size = Pt(9)

    story = [Spacer(1,35*mm), Paragraph("BÁO CÁO THỰC HÀNH", styles["TitleVI"]),
             Paragraph("DAY27 — TRACK 3<br/>HITL PR REVIEW AGENT", styles["TitleVI"]), Spacer(1,12*mm)]
    cover = [("Học viên", "Mai Quang Dũng"), ("Mã học viên", "2A202602966"),
             ("Khóa", "K4"), ("Bài thực hành", "LAB12 / Day27 Track03"), ("Ngày hoàn thành", "15/10/2026")]
    doc.add_paragraph("BÁO CÁO THỰC HÀNH", "Title")
    doc.add_paragraph("DAY27 — TRACK 3\nHITL PR REVIEW AGENT", "Title")
    for label, value in cover:
        story.append(Paragraph(f"<b>{label}:</b> {value}", styles["BodyVI"]))
        doc.add_paragraph(f"{label}: {value}")
    story.append(Spacer(1,10*mm))
    story.append(Paragraph("Repository bài nộp:<br/>https://github.com/DungMQ/K4-DAY27-MaiQuangDung-2A202602966", styles["BodyVI"]))
    doc.add_paragraph("Repository: https://github.com/DungMQ/K4-DAY27-MaiQuangDung-2A202602966")
    story.append(PageBreak())
    doc.add_page_break()

    tokens = MarkdownIt("commonmark").enable("table").parse(Path("BAO_CAO.md").read_text(encoding="utf-8"))
    start = next(i-1 for i,t in enumerate(tokens) if t.type == "inline" and t.content == "Mục lục")
    i, list_depth, references_start = start, 0, None
    while i < len(tokens):
        token = tokens[i]
        if token.type == "heading_open":
            level = min(int(token.tag[1:]),3)
            content = tokens[i+1]
            if content.content == "11.2. Tài liệu tham khảo":
                references_start = len(story)
            story.append(Paragraph(inline_pdf(content), styles[f"HeadingVI{level}"]))
            doc.add_heading(plain(content), level=level)
            i += 3
            continue
        if token.type == "table_open":
            rows, row = [], []
            i += 1
            while tokens[i].type != "table_close":
                item = tokens[i]
                if item.type == "tr_open": row = []
                elif item.type == "inline": row.append(item)
                elif item.type == "tr_close": rows.append(row)
                i += 1
            widths = [170*mm / len(rows[0])] * len(rows[0])
            table = Table([[Paragraph(inline_pdf(cell),styles["TableVI"]) for cell in row] for row in rows], colWidths=widths, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#e9f0f4")),
                ("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#b7c6cf")),
                ("VALIGN",(0,0),(-1,-1),"TOP"), ("LEFTPADDING",(0,0),(-1,-1),5),
                ("RIGHTPADDING",(0,0),(-1,-1),5), ("TOPPADDING",(0,0),(-1,-1),5), ("BOTTOMPADDING",(0,0),(-1,-1),5),
            ]))
            story.extend([table,Spacer(1,7)])
            word_table = doc.add_table(rows=0, cols=len(rows[0]))
            word_table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                cells = word_table.add_row().cells
                if row_index == 0:
                    header = OxmlElement("w:tblHeader")
                    word_table.rows[-1]._tr.get_or_add_trPr().append(header)
                for j, cell in enumerate(row):
                    cells[j].text = plain(cell)
                    for run in cells[j].paragraphs[0].runs:
                        run.font.size = Pt(9)
                        run.bold = row_index == 0
        elif token.type == "fence":
            code = "\n".join(textwrap.fill(line,width=94,replace_whitespace=False,drop_whitespace=False) if len(line)>94 else line
                             for line in token.content.rstrip().splitlines())
            story.append(Preformatted(code,code_style))
            p = doc.add_paragraph()
            run = p.add_run(token.content.rstrip())
            run.font.name, run.font.size = "Consolas", Pt(8)
        elif token.type == "paragraph_open":
            content = tokens[i+1]
            images = [child for child in content.children or [] if child.type == "image"]
            if images:
                for img in images:
                    path = img.attrGet("src")
                    with PILImage.open(path) as bitmap:
                        width, height = bitmap.size
                    scale = min(170*mm/width, 155*mm/height)
                    story.extend([Image(path,width=width*scale,height=height*scale),Spacer(1,7)])
                    doc.add_picture(path, width=Cm(width*scale/(10*mm)))
            else:
                text = inline_pdf(content)
                prefix = "• " if list_depth else ""
                story.append(Paragraph(prefix+text, styles["BodyVI"]))
                paragraph = doc.add_paragraph(plain(content), style="List Bullet" if list_depth else None)
                if references_start is not None:
                    paragraph.paragraph_format.keep_with_next = True
            i += 2
        elif token.type in {"bullet_list_open","ordered_list_open"}: list_depth += 1
        elif token.type in {"bullet_list_close","ordered_list_close"}: list_depth -= 1
        i += 1

    if references_start is not None:
        story = story[:references_start] + [KeepTogether(story[references_start:])]
        doc.paragraphs[-1].paragraph_format.keep_with_next = False
    doc.core_properties.title = "Báo cáo Day27 Track 3 — HITL PR Review Agent"
    doc.core_properties.author = "Mai Quang Dũng"
    doc.save(OUT / (NAME + ".docx"))

    def page_footer(canvas, document):
        canvas.setFont("Report",9)
        canvas.setFillColor(colors.HexColor("#4d6170"))
        canvas.drawString(20*mm,12*mm,"Mai Quang Dũng | Day27 Track 3")
        canvas.drawRightString(190*mm,12*mm,f"Trang {document.page}")

    pdf = SimpleDocTemplate(str(OUT / (NAME+".pdf")),pagesize=A4, leftMargin=20*mm,rightMargin=20*mm,
                            topMargin=18*mm,bottomMargin=20*mm,title=doc.core_properties.title,author="Mai Quang Dũng")
    pdf.build(story,onFirstPage=page_footer,onLaterPages=page_footer)
    rendered = pymupdf.open(OUT / (NAME+".pdf"))
    text = "\n".join(page.get_text() for page in rendered)
    required = ["Mai Quang Dũng","2A202602966","16","SQLite","Kết luận"]
    missing = [item for item in required if item not in text]
    if missing: raise SystemExit("Missing PDF text: " + ", ".join(missing))
    Path("scratch").mkdir(exist_ok=True)
    for index in [0,1,len(rendered)-1]:
        rendered[index].get_pixmap(matrix=pymupdf.Matrix(1.3,1.3)).save(f"scratch/report-preview-{index+1:02}.png")
    validation = {"pdf_pages":len(rendered),"pdf_images":sum(len(p.get_images()) for p in rendered),
                  "vietnamese_text_verified": True, "docx_tables":len(doc.tables),"docx_images":len(doc.inline_shapes),
                  "source_sha256": hashlib.sha256(Path("BAO_CAO.md").read_bytes()).hexdigest(),
                  "files":[str(OUT/(NAME+ext)) for ext in (".pdf",".docx")]}
    (OUT / "validation.json").write_text(json.dumps(validation,indent=2,ensure_ascii=False),encoding="utf-8")
    print(json.dumps(validation,indent=2,ensure_ascii=False))


if __name__ == "__main__":
    main()
